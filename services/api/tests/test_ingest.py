from io import BytesIO

import pytest
from docx import Document

from app.core.config import Settings
from app.services.ingest import DOI_PATTERN, ingest, ingest_pasted_text


def _settings(tmp_path) -> Settings:
    return Settings(data_dir=tmp_path, max_pdf_pages=5, max_vision_pages=2)


def test_doi_regex_extracts_common_shapes():
    text = "See https://doi.org/10.1038/nature12373 and also 10.1111/1467-8624.00164 for context."
    assert "10.1038/nature12373" in DOI_PATTERN.findall(text)
    assert "10.1111/1467-8624.00164" in DOI_PATTERN.findall(text)


def test_docx_ingest_collects_paragraphs(tmp_path):
    doc = Document()
    doc.add_heading("Sample", level=1)
    doc.add_paragraph("This references DOI 10.1038/nature12373 for completeness.")
    buffer = BytesIO()
    doc.save(buffer)
    payload = buffer.getvalue()

    result = ingest(payload, "sample.docx", _settings(tmp_path))

    assert result.detected_type == "docx"
    assert "10.1038/nature12373" in result.dois
    assert "Sample" in result.text


def test_pasted_text_ingest_extracts_doi():
    result = ingest_pasted_text("Smith (2020). Study. Nature, 581. https://doi.org/10.1038/s41586-020-2715-9")
    assert result.detected_type == "pasted_text"
    assert result.dois == ["10.1038/s41586-020-2715-9"]
    assert result.text.startswith("Smith")


def test_unsupported_type_raises(tmp_path):
    with pytest.raises(ValueError):
        ingest(b"hello", "weird.xyz", _settings(tmp_path))
