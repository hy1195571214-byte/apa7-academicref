"""Document ingestion: unify PDF / DOCX / images into text + evidence pages."""

from __future__ import annotations

import io
import logging
import mimetypes
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

import fitz  # pymupdf
from docx import Document
from PIL import Image

from ..core.config import Settings

logger = logging.getLogger(__name__)

DOI_PATTERN = re.compile(r"10\.\d{4,9}/[^\s\"'<>]+", re.IGNORECASE)
REFERENCES_HINTS = (
    "references",
    "bibliography",
    "works cited",
    "参考文献",
    "引用文献",
)


@dataclass
class IngestedDocument:
    text: str
    evidence_pages: List[int] = field(default_factory=list)
    candidate_images: List[bytes] = field(default_factory=list)
    detected_type: str = "unknown"
    dois: List[str] = field(default_factory=list)


def detect_mime(filename: str, fallback: Optional[str] = None) -> str:
    guess, _ = mimetypes.guess_type(filename)
    if guess:
        return guess
    if fallback:
        return fallback
    return "application/octet-stream"


def ingest(payload: bytes, filename: str, settings: Settings) -> IngestedDocument:
    mime = detect_mime(filename)
    lower = filename.lower()
    if lower.endswith(".pdf") or mime == "application/pdf":
        return _ingest_pdf(payload, settings)
    if lower.endswith(".docx") or mime.endswith("wordprocessingml.document"):
        return _ingest_docx(payload)
    if mime.startswith("image/") or lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        return _ingest_image(payload, mime or "image/png")
    raise ValueError(f"Unsupported file type: {mime or filename}")


def ingest_pasted_text(text: str) -> IngestedDocument:
    return IngestedDocument(
        text=text.strip(),
        evidence_pages=[],
        candidate_images=[],
        detected_type="pasted_text",
        dois=list(dict.fromkeys(DOI_PATTERN.findall(text))),
    )


def _ingest_docx(payload: bytes) -> IngestedDocument:
    with io.BytesIO(payload) as buf:
        doc = Document(buf)
    parts: List[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    text = "\n".join(parts)
    return IngestedDocument(
        text=text,
        evidence_pages=[],
        candidate_images=[],
        detected_type="docx",
        dois=list(dict.fromkeys(DOI_PATTERN.findall(text))),
    )


def _ingest_pdf(payload: bytes, settings: Settings) -> IngestedDocument:
    doc = fitz.open(stream=payload, filetype="pdf")
    try:
        total_pages = min(doc.page_count, settings.max_pdf_pages)
        page_texts: List[str] = []
        for i in range(total_pages):
            page_texts.append(doc.load_page(i).get_text("text") or "")

        joined_text = "\n".join(page_texts)
        dois = list(dict.fromkeys(DOI_PATTERN.findall(joined_text)))

        evidence_pages = _select_evidence_pages(page_texts, settings.max_vision_pages)
        candidate_images: List[bytes] = []
        text_is_thin = sum(len(t.strip()) for t in page_texts) < 200
        should_render_images = text_is_thin or not joined_text.strip()
        if should_render_images and evidence_pages:
            for page_index in evidence_pages[: settings.max_vision_pages]:
                page = doc.load_page(page_index)
                pix = page.get_pixmap(dpi=180)
                candidate_images.append(pix.tobytes("png"))

        text = _trim_text(joined_text)
        return IngestedDocument(
            text=text,
            evidence_pages=evidence_pages,
            candidate_images=candidate_images,
            detected_type="pdf",
            dois=dois,
        )
    finally:
        doc.close()


def _ingest_image(payload: bytes, mime: str) -> IngestedDocument:
    with Image.open(io.BytesIO(payload)) as img:
        img.verify()
    return IngestedDocument(
        text="",
        evidence_pages=[0],
        candidate_images=[payload],
        detected_type=f"image/{mime.split('/')[-1]}",
        dois=[],
    )


def _select_evidence_pages(page_texts: List[str], limit: int) -> List[int]:
    """Pick pages most likely to contain citation metadata."""

    if not page_texts:
        return []
    scored: List[tuple[int, int]] = []
    for idx, text in enumerate(page_texts):
        score = 0
        lower = text.lower()
        if idx == 0:
            score += 3
        if idx == len(page_texts) - 1:
            score += 1
        for hint in REFERENCES_HINTS:
            if hint in lower:
                score += 4
                break
        if DOI_PATTERN.search(text):
            score += 5
        if score > 0:
            scored.append((score, idx))
    scored.sort(key=lambda item: (-item[0], item[1]))
    ordered = [idx for _, idx in scored]
    if 0 not in ordered:
        ordered.insert(0, 0)
    deduped: List[int] = []
    for idx in ordered:
        if idx not in deduped:
            deduped.append(idx)
        if len(deduped) >= max(limit, 2):
            break
    return deduped


def _trim_text(text: str, max_chars: int = 24_000) -> str:
    if len(text) <= max_chars:
        return text
    half = max_chars // 2
    return text[:half] + "\n...\n" + text[-half:]
